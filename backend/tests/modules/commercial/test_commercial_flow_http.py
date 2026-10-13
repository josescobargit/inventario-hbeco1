"""Every result surface must agree without asking for the files a second time."""

import io
import os
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.main import app
from app.core.database import Base, get_db
from app.modules.auth.api.dependencies import get_current_user
from app.modules.catalog.infrastructure.models import Product
from app.modules.commercial.models import CommercialDocument, CommercialRecord


@pytest.fixture
def client():
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = Session(engine)
    db.add_all(
        [
            Product(
                id=uuid.uuid4(),
                sku=sku,
                name=name,
                category="Capilar",
                units_per_box=12,
                cost=Decimal("1"),
            )
            for sku, name in (("P1", "ROMERO"), ("P2", "COCO"))
        ]
    )
    db.commit()
    user = SimpleNamespace(id=uuid.uuid4(), role=SimpleNamespace(code="principal"))
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    try:
        yield TestClient(app), db
    finally:
        app.dependency_overrides.clear()
        db.close()
        engine.dispose()


def upload(http, filename, rows):
    content = "Tipo,Cadena,Fecha,OC,Factura,Codigo,Unidades,Precio unitario,Moneda,Estado,Reemplaza\n"
    content += "\n".join(",".join(map(str, row)) for row in rows)
    response = http.post(
        "/api/v1/commercial/imports",
        files={"files": (filename, content.encode(), "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_multiple_invoices_share_one_comparison_dashboard_and_weekly_report(client):
    http, db = client
    upload(
        http,
        "order.csv",
        [
            ("OC", "TIA", "2026-09-14", "12345", "", "P1", 10, 2, "USD", "activo", ""),
            ("OC", "TIA", "2026-09-14", "12345", "", "P2", 10, 3, "USD", "activo", ""),
        ],
    )
    upload(
        http,
        "invoice-1.csv",
        [
            (
                "factura",
                "TIA",
                "2026-09-15",
                "12345",
                "F1",
                "P1",
                6,
                2,
                "USD",
                "activo",
                "",
            )
        ],
    )
    upload(
        http,
        "invoice-2.csv",
        [
            (
                "factura",
                "TIA",
                "2026-09-16",
                "12345",
                "F2",
                "P1",
                6,
                2,
                "USD",
                "activo",
                "",
            )
        ],
    )
    upload(
        http,
        "void.csv",
        [
            (
                "factura",
                "TIA",
                "2026-09-16",
                "12345",
                "F3",
                "P2",
                10,
                3,
                "USD",
                "anulada",
                "",
            )
        ],
    )
    params = "date_from=2026-09-14&date_to=2026-09-20"
    workflow = http.get(f"/api/v1/commercial/workflow?{params}").json()
    dashboard = http.get(f"/api/v1/commercial/dashboard?{params}").json()
    report_response = http.get(f"/api/v1/commercial/weekly-report?{params}")
    assert report_response.status_code == 200, report_response.text
    report = report_response.json()
    assert len(workflow["orders"][0]["invoices"]) == 2
    assert len(workflow["orders"]) == 1
    assert workflow["comparisons"][0]["invoiced"] == 12
    for result in (dashboard, report):
        summary = result["summary"]
        assert summary["ordered"] == 20
        assert summary["invoiced"] == 12
        assert summary["missing"] == 10
        assert summary["fulfillment"] == 50
        assert summary["unbilled_amount"] == 30
        assert result["chains"][0]["missing"] == 10
    assert report["period"] == {
        "date_from": "2026-09-14",
        "date_to": "2026-09-20",
        "basis": report["basis"],
    }
    assert {r["sku"] for r in report["critical_products"]["not_invoiced"]} == {"P2"}
    assert any(a["code"] == "excluded" for a in report["alerts"])
    assert len(list(db.scalars(select(CommercialDocument)))) == 4
    assert len(list(db.scalars(select(CommercialRecord)))) == 5
    assert all(not row.corrections for row in db.scalars(select(CommercialRecord)))
    export = http.get(f"/api/v1/commercial/export?{params}&format=xlsx")
    assert export.status_code == 200, export.text
    workbook = load_workbook(io.BytesIO(export.content))
    sheet = workbook["OC vs factura"]
    data = [
        dict(zip([cell.value for cell in sheet[1]], row))
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert sum(row["FALTANTE"] for row in data) == 10
    assert sum(row["DINERO NO FACTURADO"] for row in data) == 30


def test_period_selection_and_default_week_are_explicit(client):
    http, _ = client
    upload(
        http,
        "order.csv",
        [("OC", "TIA", "2026-09-14", "12345", "", "P1", 10, 2, "USD", "activo", "")],
    )
    future = http.get(
        "/api/v1/commercial/dashboard?date_from=2026-10-01&date_to=2026-10-07"
    ).json()
    assert future["rows"] == []
    assert (
        http.get(
            "/api/v1/commercial/weekly-report?date_from=2026-10-20&date_to=2026-10-01"
        ).status_code
        == 422
    )
    report = http.get("/api/v1/commercial/weekly-report").json()
    today = datetime.now(ZoneInfo("America/Guayaquil")).date()
    assert date.fromisoformat(report["period"]["date_from"]) == today - timedelta(
        days=today.weekday()
    )
    assert date.fromisoformat(report["period"]["date_to"]) == today


def test_ambiguous_invoice_keeps_totals_pending_and_exports_work(client):
    http, _ = client
    for number in ("12345", "67890"):
        upload(
            http,
            f"{number}.csv",
            [("OC", "TIA", "2026-09-14", number, "", "P1", 10, 2, "USD", "activo", "")],
        )
    upload(
        http,
        "invoice.csv",
        [("factura", "TIA", "2026-09-15", "", "F1", "P1", 6, 2, "USD", "activo", "")],
    )
    params = "date_from=2026-09-14&date_to=2026-09-20"
    data = http.get(f"/api/v1/commercial/dashboard?{params}").json()
    assert data["summary"]["invoiced"] is None
    assert data["summary"]["review_documents"] == 1
    for format_ in ("xlsx", "pdf"):
        response = http.get(
            f"/api/v1/commercial/export?{params}&format={format_}&view=missing"
        )
        assert response.status_code == 200, response.text


@pytest.mark.skipif(
    not os.environ.get("REAL_INVOICE_DIR") or not os.environ.get("REAL_OC_ZIP"),
    reason="Requires the user's unmodified real OC and invoice files",
)
def test_real_danec_from_upload_through_weekly_report(client):
    import zipfile
    from scripts.seed_products import load_rows

    http, db = client
    for values in load_rows():
        db.add(Product(id=uuid.uuid4(), **values))
    db.commit()
    invoice_path = Path(os.environ["REAL_INVOICE_DIR"]) / "FAC 001-001-000001016.pdf"
    with zipfile.ZipFile(os.environ["REAL_OC_ZIP"]) as archive:
        order_bytes = archive.read("OC 26003023.pdf")
    for name, content in (
        ("documento-oc.pdf", order_bytes),
        ("documento-factura.pdf", invoice_path.read_bytes()),
    ):
        response = http.post(
            "/api/v1/commercial/imports",
            files={"files": (name, content, "application/pdf")},
        )
        assert response.status_code == 200, response.text
        assert response.json()["files"][0]["status"] == "processed"
    params = "date_from=2026-09-14&date_to=2026-09-20&chain=INDUSTRIAL%20DANEC"
    workflow = http.get(f"/api/v1/commercial/workflow?{params}").json()
    assert len(workflow["orders"]) == len(workflow["invoices"]) == 1
    assert workflow["orders"][0]["number"] == "26003023"
    assert workflow["invoices"][0]["number"] == "001-001-000001016"
    assert workflow["invoices"][0]["linked_order_id"] == workflow["orders"][0]["id"]
    assert workflow["invoices"][0]["status"] == "OK"
    assert len(workflow["comparisons"]) == 4
    assert all(row["status"] == "COMPLETO" for row in workflow["comparisons"])
    dashboard = http.get(f"/api/v1/commercial/dashboard?{params}").json()
    report = http.get(f"/api/v1/commercial/weekly-report?{params}").json()
    for result in (dashboard, report):
        assert result["summary"]["ordered"] == result["summary"]["invoiced"] == 1248
        assert result["summary"]["fulfillment"] == 100
        assert result["summary"]["missing"] == 0
        assert result["summary"]["invoiced_amount"] == 5218.08
    assert len(list(db.scalars(select(CommercialDocument)))) == 2


@pytest.mark.skipif(
    not os.environ.get("REAL_INVOICE_DIR"),
    reason="Requires the user's real OC/invoice documents outside the repository",
)
@pytest.mark.parametrize(
    "chain,order_file,number,invoice_numbers",
    [
        ("CORPORACIÓN FAVORITA", "OC 14 SEPT - 1.pdf", "100625388996", [1009]),
        ("CORPORACIÓN EL ROSADO", "OC 2-JUL.pdf", "4617985918", [707]),
        ("TUTI", "OC 4500360076.PDF", "4500360076", [915, 916]),
        ("TIA", "OC TIA 3000949552 .pdf", "3000949552", [820]),
        ("LIRIS DEL PORTAL", "OC_005884445.pdf", "OC_005884445", [913]),
        ("FARCOMED", "orden-compra-6974083.pdf", "6974083", [978]),
        ("GERARDO ORTIZ", "CP143221.pdf", "5601886574", [1107]),
    ],
)
def test_real_chain_pair_keeps_all_invoices_and_same_result_surfaces(
    client, chain, order_file, number, invoice_numbers
):
    from scripts.seed_products import load_rows

    http, db = client
    for values in load_rows():
        db.add(Product(id=uuid.uuid4(), **values))
    db.commit()
    directory = Path(os.environ["REAL_INVOICE_DIR"])
    paths = [
        directory / order_file,
        *(directory / f"FAC 001-001-{n:09d}.pdf" for n in invoice_numbers),
    ]
    for index, path in enumerate(paths):
        response = http.post(
            "/api/v1/commercial/imports",
            files={
                "files": (f"original-{index}.pdf", path.read_bytes(), "application/pdf")
            },
        )
        assert response.status_code == 200, response.text
        assert response.json()["files"][0]["status"] == "processed"
    workflow = http.get("/api/v1/commercial/workflow").json()
    assert len(workflow["orders"]) == (2 if chain == "CORPORACIÓN EL ROSADO" else 1)
    matched_orders = [o for o in workflow["orders"] if o["number"] == number]
    assert len(matched_orders) == 1
    order = matched_orders[0]
    assert order["number"] == number
    assert order["chain"] == chain
    expected_products = {
        "CORPORACIÓN FAVORITA": 13,
        "CORPORACIÓN EL ROSADO": 12,
        "TUTI": 3,
        "TIA": 4,
        "LIRIS DEL PORTAL": 3,
        "FARCOMED": 2,
        "GERARDO ORTIZ": 6,
    }
    assert len(order["lines"]) == expected_products[chain]
    if chain == "TIA":
        assert order["quantity"] == 804
        assert all("Derechos Reservados" not in line["name"] for line in order["lines"])
    assert len(workflow["invoices"]) == len(invoice_numbers)
    assert all(
        invoice["linked_order_id"] == order["id"] for invoice in workflow["invoices"]
    ), [(i["number"], i["reference"], i["issues"]) for i in workflow["invoices"]]
    assert len(order["invoices"]) == len(invoice_numbers)
    dashboard = http.get("/api/v1/commercial/dashboard").json()
    report = http.get(
        "/api/v1/commercial/weekly-report?date_from=2026-01-01&date_to=2026-12-31"
    ).json()
    expected_rows = [
        (r["code"], r["ordered"], r["invoiced"], r["missing"], r["status"])
        for r in workflow["comparisons"]
    ]
    assert [
        (r["sku"], r["ordered"], r["invoiced"], r["missing"], r["status"])
        for r in dashboard["rows"]
    ] == expected_rows
    for key in (
        "ordered",
        "invoiced",
        "missing",
        "fulfillment",
        "ordered_amount",
        "invoiced_amount",
        "unbilled_amount",
    ):
        assert dashboard["summary"][key] == report["summary"][key]
    assert len(list(db.scalars(select(CommercialDocument)))) == len(paths)
