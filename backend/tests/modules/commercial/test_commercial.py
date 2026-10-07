import io
import uuid
import zipfile
from decimal import Decimal
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.main import app
from app.core.database import Base, get_db
from app.modules.auth.api.dependencies import get_current_user
from app.modules.catalog.infrastructure.models import Product
from app.modules.commercial.ingestion import parse_file, unpack
from app.modules.commercial.models import CommercialRecord
from app.modules.commercial.normalization import normalize, number
from app.modules.commercial.service import analyze
from app.modules.audit.infrastructure.models import AuditLog


@pytest.fixture
def product():
    return Product(
        id=uuid.uuid4(),
        sku="SH001",
        name="SHAMPOO ROMERO 500 ML",
        category="Capilar",
        barcode="7862133169541",
        units_per_box=12,
        cost=Decimal("1"),
    )


def normalized(product, **raw):
    defaults = dict(
        kind="order",
        chain="Cadena A",
        date="2026-10-01",
        product_id=str(product.id),
        order_number="OC-1",
        total_units=12,
        unit_price=2,
        currency="USD",
    )
    defaults.update(raw)
    result = normalize(defaults, [product], {}, [])
    result.update(
        id=str(uuid.uuid4()),
        document_id=str(uuid.uuid4()),
        filename="documento.csv",
        origin="import",
        source_url="/source",
        source_location="Fila 2",
        source=defaults,
        revision=1,
    )
    return result


def test_conversion_and_price_are_auditable(product):
    row = normalized(product, quantity=2, unit="cajas", total_units=None, unit_price=24)
    assert row["units"] == 24
    assert row["unit_price"] == "2"
    assert row["amount"] == "48.00"
    assert "2 cajas × 12" in row["conversion"]
    assert row["issues"] == []


def test_ambiguous_number_never_silently_changes_scale():
    for value in ("1.000", "1,000", "NaN", "Infinity"):
        with pytest.raises(ValueError):
            number(value)
    assert number("1.234,50") == Decimal("1234.50")
    assert number("1,234.50") == Decimal("1234.50")
    assert number(1.123) == Decimal("1.123")


def test_conflicting_units_and_ean_require_review(product):
    row = normalized(product, total_units=10, boxes=2)
    assert any("difieren" in reason for reason in row["issues"])
    row = normalized(product, product_id=None, code="SH001", ean="9999999999999")
    assert row["product_id"] is None
    assert row["issues"]


def test_family_equivalence_preserves_size_and_type(product):
    row = normalized(
        product, product_id=None, product_name="SHAMPOO CRECIMIENTO 500 ML"
    )
    assert row["product_id"] == str(product.id)
    assert (
        normalized(product, product_id=None, product_name="SHAMPOO CRECIMIENTO 370 ML")[
            "product_id"
        ]
        is None
    )
    assert (
        normalized(
            product, product_id=None, product_name="ACONDICIONADOR CRECIMIENTO 500 ML"
        )["product_id"]
        is None
    )


def test_multiple_invoices_sum_and_keep_each_source(product):
    order = normalized(product, total_units=24)
    a = normalized(product, kind="invoice", invoice_number="F1", total_units=8)
    b = normalized(product, kind="invoice", invoice_number="F2", total_units=10)
    data = analyze([order, a, b])
    row = data["rows"][0]
    assert (row["invoiced"], row["missing"], row["status"]) == (18, 6, "PARCIAL")
    assert row["unbilled_amount"] == 12
    assert {r["invoice_number"] for r in row["invoices"]} == {"F1", "F2"}
    assert data["summary"]["fulfillment"] == 75


def test_cancellation_and_duplicates_do_not_inflate_delivery(product):
    order = normalized(product)
    void = normalized(product, kind="invoice", invoice_number="VOID", status="anulada")
    a = normalized(product, kind="invoice", invoice_number="DUP")
    b = normalized(product, kind="invoice", invoice_number="DUP")
    result = analyze([order, void, a, b])
    assert result["summary"]["invoiced"] == 0
    assert result["rows"][0]["status"] == "REVISAR"
    assert any(a["code"] == "duplicate" for a in result["alerts"])
    assert any(a["code"] == "excluded" for a in result["alerts"])


def test_missing_prices_never_use_product_cost(product):
    result = analyze([normalized(product, unit_price=None)])
    assert result["summary"]["ordered_amount"] is None
    assert result["summary"]["unbilled_amount"] is None
    assert result["summary"]["missing"] == 12


def test_excess_of_one_product_never_compensates_another(product):
    second = Product(
        id=uuid.uuid4(),
        sku="SH002",
        name="COCO",
        category="Capilar",
        units_per_box=12,
        cost=Decimal(1),
    )
    result = analyze(
        [
            normalized(product, total_units=10),
            normalized(second, total_units=10),
            normalized(product, kind="invoice", invoice_number="F1", total_units=20),
        ]
    )
    # Same OC, two original rows within one document.
    # Give both OC rows the same document id rather than making duplicate documents.
    inputs = [
        normalized(product, total_units=10),
        normalized(second, total_units=10),
        normalized(product, kind="invoice", invoice_number="F1", total_units=20),
    ]
    inputs[1]["document_id"] = inputs[0]["document_id"]
    result = analyze(inputs)
    assert result["summary"]["fulfillment"] == 50
    assert result["summary"]["missing"] == 10


def test_duplicate_detection_happens_before_date_filters(product):
    order = normalized(product)
    a = normalized(product, kind="invoice", invoice_number="F1", date="2026-09-30")
    b = normalized(product, kind="invoice", invoice_number="F1", date="2026-10-02")
    result = analyze([order, a, b], start="2026-10-01", end="2026-10-31")
    assert result["summary"]["invoiced"] == 0


def test_sell_out_absence_is_not_zero_and_coverage_must_match(product):
    a = normalized(product, kind="sell_in", date="2026-10-01")
    b = normalized(product, kind="sell_out", date="2026-10-02", total_units=2)
    result = analyze([a, b])
    assert result["sell_comparison"][0]["difference"] is None
    assert "Cobertura incompleta" in result["sell_comparison"][0]["hypothesis"]
    result = analyze([a])
    assert result["summary"]["sell_out_units"] is None


def test_zip_path_traversal_is_rejected():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("../outside.csv", "data")
    with pytest.raises(ValueError, match="ruta"):
        unpack("upload.zip", buffer.getvalue())


def test_excel_multiple_sheets_and_formulas_are_not_assumed():
    book = Workbook()
    for sheet in [book.active, book.create_sheet("Octubre")]:
        sheet.append(["Producto", "Unidades", "Subtotal"])
        sheet.append(["Shampoo", 12, "=12*2"])
    buffer = io.BytesIO()
    book.save(buffer)
    _, _, _, rows, _ = parse_file("report.xlsx", buffer.getvalue())
    assert len(rows) == 2
    assert all("fórmulas" in r["extraction_issue"] for r in rows)


@pytest.fixture
def client(product):
    engine = create_engine(
        "sqlite+pysqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    db = Session(engine)
    db.add(product)
    db.commit()
    user = SimpleNamespace(id=uuid.uuid4(), role=SimpleNamespace(code="principal"))
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    yield TestClient(app), db
    app.dependency_overrides.clear()
    db.close()
    engine.dispose()


def csv_bytes(kind="OC", number="OC-1", quantity=24):
    return f"Tipo,Cadena,Fecha,OC,Factura,Codigo,Unidades,Precio unitario,Moneda\n{kind},Cadena A,2026-10-01,OC-1,{number if kind == 'factura' else ''},SH001,{quantity},2,USD\n".encode()


def test_end_to_end_import_review_export_preserves_history(client):
    http, db = client
    response = http.post(
        "/api/v1/commercial/imports",
        files=[
            ("files", ("order.csv", csv_bytes(), "text/csv")),
            ("files", ("invoice.csv", csv_bytes("factura", "F1", 12), "text/csv")),
        ],
    )
    assert response.status_code == 200, response.text
    assert all(f["status"] == "processed" for f in response.json()["files"])
    data = http.get("/api/v1/commercial/dashboard").json()
    assert data["summary"]["ordered"] == 24
    assert data["summary"]["invoiced"] == 12
    assert data["summary"]["unbilled_amount"] == 24
    duplicate = http.post(
        "/api/v1/commercial/imports",
        files={"files": ("order.csv", csv_bytes(), "text/csv")},
    ).json()
    assert duplicate["files"][0]["status"] == "duplicate"
    record = next(r for r in data["records"] if r["kind"] == "invoice")
    original = db.get(CommercialRecord, uuid.UUID(record["id"])).source.copy()
    change = {
        "revision": 1,
        "reason": "Anulación verificada en documento original",
        "values": {"status": "anulada"},
    }
    assert (
        http.patch(
            f"/api/v1/commercial/records/{record['id']}", json=change
        ).status_code
        == 200
    )
    assert (
        http.patch(
            f"/api/v1/commercial/records/{record['id']}", json=change
        ).status_code
        == 409
    )
    assert db.get(CommercialRecord, uuid.UUID(record["id"])).source == original
    assert (
        db.scalar(select(AuditLog).where(AuditLog.action == "commercial.correct"))
        is not None
    )
    assert http.get("/api/v1/commercial/dashboard").json()["summary"]["invoiced"] == 0
    xlsx = http.get("/api/v1/commercial/export?format=xlsx")
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book["OC vs factura"]["E2"].value == 24
    assert book["OC vs factura"]["E2"].data_type == "n"
    pdf = http.get("/api/v1/commercial/export?format=pdf")
    assert pdf.content.startswith(b"%PDF")
    source = http.get(record["source_url"].join(["/api/v1", ""]))
    assert source.content == csv_bytes("factura", "F1", 12)
    assert len(http.get("/api/v1/commercial/documents").json()) == 2
    answer = http.post(
        "/api/v1/commercial/assistant",
        json={"question": "Hazme un resumen para gerencia"},
    ).json()
    assert any("24 unidades" in s for s in answer["answer"])
    assert answer["sources"]


def test_authentication_required():
    http = TestClient(app)
    assert http.get("/api/v1/commercial/dashboard").status_code == 401
    assert (
        http.post(
            "/api/v1/commercial/assistant", json={"question": "Resumen"}
        ).status_code
        == 401
    )


def test_wrong_date_range_is_rejected(client):
    http, _ = client
    response = http.get(
        "/api/v1/commercial/dashboard?date_from=2026-10-05&date_to=2026-10-01"
    )
    assert response.status_code == 422


def test_price_per_unit_is_not_divided_again_when_boxes_are_also_present(product):
    row = normalized(product, quantity=24, unit="unidades", boxes=2, unit_price=2)
    assert row["issues"]  # Original explicit total is 12, inconsistent with 24.
    row = normalized(
        product, total_units=24, quantity=24, unit="unidades", boxes=2, unit_price=2
    )
    assert not row["issues"]
    assert row["amount"] == "48.00"
    assert row["unit_price"] == "2"


def test_mismatched_subtotal_is_not_reported_as_confirmed_money(product):
    row = normalized(product, amount=10)
    assert row["amount"] is None
    assert row["monetary_issues"]


def test_replacement_without_void_evidence_blocks_both_invoices(product):
    old = normalized(product, kind="invoice", invoice_number="OLD")
    new = normalized(product, kind="invoice", invoice_number="NEW", replaces="OLD")
    result = analyze([normalized(product), old, new])
    assert result["summary"]["invoiced"] == 0
    assert any(a["code"] == "replacement" for a in result["alerts"])


def test_header_metadata_and_sellout_decline_are_traceable(product):
    content = b"Sell Out;;;;;\nCadena;Cadena A;;;;\nMoneda;USD;;;;\nFecha;2026-10-01;;;;\nProducto;Unidades;Subtotal;;;\nSHAMPOO ROMERO 500 ML;5;10;;;\n"
    _, _, _, rows, _ = parse_file("report.csv", content)
    assert rows[0]["kind"] == "sell_out"
    assert rows[0]["chain"] == "Cadena A"
    assert rows[0]["currency"] == "USD"
    a = normalized(product, kind="sell_out", date="2026-09-01", total_units=100)
    b = normalized(product, kind="sell_out", date="2026-10-01", total_units=60)
    data = analyze([a, b])
    assert data["sales_movements"][0]["variation"] == -40
    assert any(
        a["code"] == "sell_out_decline" and "cobertura" in a["message"]
        for a in data["alerts"]
    )


def test_document_header_only_completes_missing_cells(client):
    http, db = client
    payload = b"Tipo,Cadena,Fecha,Codigo,Unidades,Subtotal,Moneda\nSell Out,,2026-10-01,SH001,2,4,\nSell Out,Otra cadena,2026-10-02,SH001,3,6,USD\n"
    result = http.post(
        "/api/v1/commercial/imports",
        files={"files": ("sellout.csv", payload, "text/csv")},
    ).json()
    doc_id = result["files"][0]["document_id"]
    response = http.patch(
        f"/api/v1/commercial/documents/{doc_id}/missing-header",
        json={
            "reason": "Cabecera del original revisada",
            "values": {"chain": "Cadena A", "currency": "USD"},
        },
    )
    assert response.status_code == 200
    assert response.json()["updated_records"] == 1
    source = list(
        db.scalars(select(CommercialRecord).order_by(CommercialRecord.row_number))
    )
    assert source[0].corrections["chain"] == "Cadena A"
    assert source[1].corrections == {}


def test_pdf_and_image_use_existing_extraction_but_require_review(monkeypatch):
    import fitz
    from PIL import Image
    from app.modules.purchase_orders.domain import document_extraction

    extracted = document_extraction.ExtractedDocument(
        text="CADENA: Cadena A\nORDEN DE COMPRA: OC-100\nFECHA: 2026-10-01",
        method="ocr",
        page_count=1,
        table_rows=(
            {
                "description": "SHAMPOO ROMERO 500 ML",
                "quantity": 2,
                "original_unit_type": "boxes",
                "page": 1,
                "raw": "2 cajas",
            },
        ),
    )
    monkeypatch.setattr(
        document_extraction, "extract_document", lambda *args: extracted
    )
    pdf = fitz.open()
    pdf.new_page()
    payload = pdf.tobytes()
    pdf.close()
    for name, content in [("order.pdf", payload)]:
        _, _, text, rows, _ = parse_file(name, content)
        assert text == extracted.text
        assert rows[0]["quantity"] == 2
        assert rows[0]["extraction_issue"]
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(image, format="PNG")
    assert parse_file("photo.png", image.getvalue())[3][0]["extraction_issue"]


def test_ai_selects_only_validated_queries_without_business_figures(monkeypatch):
    import json
    from datetime import date
    from app.modules.commercial import assistant

    settings = SimpleNamespace(
        commercial_ai_enabled=True,
        commercial_ai_model="configured-model",
        openai_api_key="test-key",
    )
    monkeypatch.setattr(assistant, "get_settings", lambda: settings)
    sent = {}

    def fake_open(request, timeout):
        sent.update(json.loads(request.data))
        return io.BytesIO(
            json.dumps(
                {
                    "output": [
                        {
                            "type": "function_call",
                            "name": "consult_commercial_data",
                            "arguments": json.dumps(
                                {
                                    "intent": "missing",
                                    "chain": None,
                                    "date_from": None,
                                    "date_to": None,
                                }
                            ),
                        }
                    ]
                }
            ).encode()
        )

    monkeypatch.setattr(assistant, "urlopen", fake_open)
    plan, mode = assistant.plan_question("Qué quedó pendiente", {}, date(2026, 10, 5))
    assert plan.intent == "missing"
    assert mode == "IA con consultas verificadas"
    assert sent["store"] is False
    assert sent["tools"][0]["strict"] is True
    assert "business_data" not in sent
    monkeypatch.setattr(
        assistant, "urlopen", lambda *a, **k: io.BytesIO(b'{"output":[]}')
    )
    assert assistant.plan_question("Resumen", {}, date(2026, 10, 5))[0] is None


def test_pack_1000ml_matches_two_500ml_components_only_with_same_identity(product):
    product.name = "PACK SH + AC ANA ELIXIR ROMERO 500 ML"
    row = normalized(
        product, product_id=None, product_name="PACK ANA ELIXIR CRECIMIENTO 1000 ML"
    )
    assert row["product_id"] == str(product.id)
    row = normalized(
        product,
        product_id=None,
        product_name="PACK SHAMPOO 500 ML + ACONDICIONADOR 500 ML ANA ELIXIR ROMERO",
    )
    assert row["product_id"] == str(product.id)
    assert (
        normalized(product, product_id=None, product_name="PACK 1000 ML")["product_id"]
        is None
    )
    assert (
        normalized(
            product, product_id=None, product_name="PACK ANA ELIXIR COCO 1000 ML"
        )["product_id"]
        is None
    )
    assert (
        normalized(
            product, product_id=None, product_name="PACK ANA ELIXIR ROMERO 740 ML"
        )["product_id"]
        is None
    )


def test_document_correction_updates_all_rows_atomically_and_preserves_sources(client):
    http, db = client
    payload = b"Tipo,Cadena,Fecha,OC,Codigo,Unidades,Moneda\nOC,,fecha incorrecta,OC-1,SH001,2,USD\nOC,,fecha incorrecta,OC-1,SH001,3,USD\n"
    result = http.post(
        "/api/v1/commercial/imports",
        files={"files": ("review.csv", payload, "text/csv")},
    ).json()
    document_id = result["files"][0]["document_id"]
    rows = list(
        db.scalars(select(CommercialRecord).order_by(CommercialRecord.row_number))
    )
    originals = [dict(r.source) for r in rows]
    body = {
        "reason": "Cabecera comprobada con documento original",
        "values": {"chain": "TIA", "date": "2026-09-08"},
        "revisions": {str(r.id): r.revision for r in rows},
    }
    url = f"/api/v1/commercial/documents/{document_id}/header"
    invalid = {**body, "values": {"date": "2026-02-30"}}
    assert http.patch(url, json=invalid).status_code == 422
    assert http.patch(url, json=body).json() == {"updated_records": 2}
    db.expire_all()
    assert all(r.corrections == body["values"] and r.revision == 2 for r in rows)
    assert [r.source for r in rows] == originals
    assert http.patch(url, json=body).status_code == 409
    assert all(r.revision == 2 for r in rows)
    data = http.get("/api/v1/commercial/dashboard").json()
    assert all(
        r["chain"] == "TIA" and r["date"] == "2026-09-08" for r in data["records"]
    )
    assert (
        len(
            list(
                db.scalars(
                    select(AuditLog).where(
                        AuditLog.action == "commercial.correct_document_header"
                    )
                )
            )
        )
        == 2
    )


def test_duplicate_copy_can_be_resolved_for_whole_document_without_deleting(client):
    http, db = client
    payload = csv_bytes()
    imported = http.post(
        "/api/v1/commercial/imports",
        files=[
            ("files", ("original.csv", payload, "text/csv")),
            ("files", ("copy.csv", payload + b"\n", "text/csv")),
        ],
    ).json()["files"]
    copy_id = imported[1]["document_id"]
    data = http.get("/api/v1/commercial/dashboard").json()
    copy_rows = [r for r in data["records"] if r["document_id"] == copy_id]
    assert any("duplicado" in issue for r in copy_rows for issue in r["issues"])
    result = http.patch(
        f"/api/v1/commercial/documents/{copy_id}/header",
        json={
            "values": {"status": "excluido"},
            "reason": "Se conserva el original y se excluye esta copia",
            "revisions": {r["id"]: r["revision"] for r in copy_rows},
        },
    )
    assert result.status_code == 200
    data = http.get("/api/v1/commercial/dashboard").json()
    assert len(data["records"]) == 2
    assert data["summary"]["ordered"] == 24
    assert not any(r["issues"] for r in data["records"])
    assert len(list(db.scalars(select(CommercialRecord)))) == 2


def test_recognition_is_independent_from_catalog_and_financial_valuation(client):
    http, _ = client
    payload = b"Tipo,Cadena,Fecha,OC,Producto,Cantidad,Unidad\nOC,INDUSTRIAL DANEC,2026-09-14,26003023,ANA REGEN SHAM BOT400ML 12U,360,unidades\n"
    http.post(
        "/api/v1/commercial/imports",
        files={"files": ("danec.csv", payload, "text/csv")},
    )
    recognized = http.get("/api/v1/commercial/processing").json()["records"]
    assert recognized[0]["issues"] == []
    assert recognized[0]["units"] == 360
    assert recognized[0]["chain"] == "INDUSTRIAL DANEC"
    assert any(
        "sin equivalencia" in issue
        for issue in http.get("/api/v1/commercial/dashboard").json()["records"][0][
            "issues"
        ]
    )


def test_reupload_reprocesses_stale_extraction_preserving_original_and_history(client):
    import hashlib
    from app.modules.commercial.models import CommercialDocument
    from app.modules.commercial.processing import PARSER_VERSION

    http, db = client
    payload = b"Tipo,Cadena,Fecha,OC,Producto,Cantidad,Unidad\nOC,INDUSTRIAL DANEC,2026-09-14,26003023,Producto uno,360,UN\nOC,INDUSTRIAL DANEC,2026-09-14,26003023,Producto dos,480,UN\nOC,INDUSTRIAL DANEC,2026-09-14,26003023,Producto tres,168,UN\nOC,INDUSTRIAL DANEC,2026-09-14,26003023,Producto cuatro,240,UN\n"
    doc = CommercialDocument(
        filename="OC 26003023.csv",
        sha256=hashlib.sha256(payload).hexdigest(),
        content_type="text/csv",
        content=payload,
        method="csv",
        warnings=[],
        created_by_user_id=uuid.uuid4(),
    )
    db.add(doc)
    db.flush()
    old = CommercialRecord(
        document_id=doc.id,
        row_number=1,
        source={"product_name": "Extracción incompleta"},
        corrections={},
    )
    db.add(old)
    db.commit()
    result = http.post(
        "/api/v1/commercial/imports",
        files={"files": (doc.filename, payload, "text/csv")},
    ).json()["files"][0]
    assert result["document_id"] == str(doc.id)
    assert result["status"] == "processed" and result["rows"] == 4
    db.expire_all()
    assert doc.content == payload and doc.parser_version == PARSER_VERSION
    assert old.superseded and old.source == {"product_name": "Extracción incompleta"}
    data = http.get("/api/v1/commercial/processing").json()["records"]
    assert len(data) == 4
    assert all(
        r["issues"] == []
        and r["order_number"] == "26003023"
        and r["date"] == "2026-09-14"
        for r in data
    )
    assert len(list(db.scalars(select(CommercialRecord)))) == 5
    assert (
        db.scalar(select(AuditLog).where(AuditLog.action == "commercial.reprocess"))
        is not None
    )
    again = http.post(
        "/api/v1/commercial/imports",
        files={"files": (doc.filename, payload, "text/csv")},
    ).json()["files"][0]
    assert again["status"] == "duplicate"
    assert len(list(db.scalars(select(CommercialRecord)))) == 5


def test_danec_header_with_numero_label_and_two_digit_order_date():
    from app.modules.purchase_orders.domain.order_profiles import recognize_known_order

    for company in ("INDUSTRIAL DANEC S A", "INDUSTRIAL DANEC S.A."):
        header = recognize_known_order(
            f"{company}\n1790040968001\nORDEN DE COMPRA Nº 26003023\nFecha de Orden: 14/09/26\nFECHA DE ENTREGA: 24/09/26"
        )
        assert header["chain_name"] == "INDUSTRIAL DANEC"
        assert header["order_number"] == "26003023"
        assert header["order_date"] == "2026-09-14"
        assert header["status"] == "OK"


def test_danec_details_stop_at_subtotal():
    import fitz
    from app.modules.purchase_orders.domain.table_extraction import (
        _extract_known_text_layouts,
    )

    text = (
        "INDUSTRIAL DANEC S A\n"
        + "\n".join(
            f"{qty}.00 UN 128341028000{i} PRODUCTO {i} - 2.8100 2INVCIVA 6 100.00"
            for i, qty in enumerate((360, 480, 168, 240), 1)
        )
        + "\nSubtotal: 5218.08\n5.00 UN 1283410280099 OBSERVACIONES - 2.8100 2INVCIVA 6 100.00"
    )
    with fitz.open() as pdf:
        page = pdf.new_page()
        rows = _extract_known_text_layouts(page, 1, text)
    assert len(rows) == 4
    assert [r["quantity"] for r in rows] == [360, 480, 168, 240]


def test_ambiguous_units_are_a_real_recognition_exception():
    from app.modules.commercial.processing import recognition_evidence

    raw = {
        "kind": "order",
        "chain": "LIRIS DEL PORTAL",
        "order_number": "OC-1",
        "date": "2026-10-05",
        "product_name": "Shampoo",
        "quantity": 1,
        "unit": "ambiguous",
    }
    record = SimpleNamespace(
        id="r", row_number=1, source=raw, corrections={}, revision=1
    )
    document = SimpleNamespace(id="d", filename="doc.pdf", method="pdf_text")
    assert recognition_evidence(record, document)["issues"] == [
        "Unidad de cantidad no identificada"
    ]


def test_workflow_uses_saved_documents_without_reupload_and_preserves_sources(client):
    http, db = client
    for filename, content in [
        ("order.csv", csv_bytes()),
        ("invoice.csv", csv_bytes("factura", "F1", 12)),
    ]:
        assert (
            http.post(
                "/api/v1/commercial/imports",
                files={"files": (filename, content, "text/csv")},
            ).status_code
            == 200
        )
    response = http.get("/api/v1/commercial/workflow")
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["summary"]["linked"] == 1
    assert data["orders"][0]["invoices"][0]["number"] == "F1"
    assert data["comparisons"][0]["ordered"] == 24
    assert data["comparisons"][0]["invoiced"] == 12
    assert data["comparisons"][0]["status"] == "PARCIAL"
    assert data["comparisons"][0]["unbilled_amount"] == 24
    assert len(list(db.scalars(select(CommercialRecord)))) == 2
    assert all(not r.corrections for r in db.scalars(select(CommercialRecord)))


def test_invoice_pdf_extracts_header_all_products_and_prices_before_footer():
    import fitz

    pdf = fitz.open()
    page = pdf.new_page(width=650, height=700)
    page.insert_text((30, 30), "FACTURA No. 001-002-000000123")
    page.insert_text((30, 50), "Fecha Emision: 16/09/2026")
    page.insert_text((30, 70), "Cliente: INDUSTRIAL DANEC S A 1790040968001")
    page.insert_text((30, 90), "Orden de Compra: 26003023")
    xs = [30, 100, 300, 370, 490, 610]
    ys = [120, 150, 180, 210, 240]
    for x in xs:
        page.draw_line((x, ys[0]), (x, ys[-1]))
    for y in ys:
        page.draw_line((xs[0], y), (xs[-1], y))
    rows = [
        ["Codigo", "Descripcion", "Unidades", "Precio unitario", "Importe"],
        ["P1", "SHAMPOO ROMERO", "5", "2.00", "10.00"],
        ["P2", "ACONDICIONADOR", "3", "4.00", "12.00"],
        ["", "Subtotal", "", "", "22.00"],
    ]
    for y, cells in zip(ys, rows):
        for x, value in zip(xs, cells):
            page.insert_text((x + 4, y + 19), value, fontsize=9)
    page.insert_text((30, 280), "TOTAL FACTURA 25.30")
    content = pdf.tobytes()
    pdf.close()
    _, _, _, records, _ = parse_file("desconocido.pdf", content)
    assert len(records) == 2, records
    assert all(r["chain"] == "INDUSTRIAL DANEC" for r in records)
    assert all(
        r["kind"] == "invoice" and r["invoice_number"] == "001-002-000000123"
        for r in records
    )
    assert all(
        r["order_number"] == "26003023" and r["date"] == "2026-09-16" for r in records
    )
    assert [r["total_units"] for r in records] == ["5", "3"]
    assert [r["unit_price"] for r in records] == ["2.00", "4.00"]
    assert [r["amount"] for r in records] == ["10.00", "12.00"]


def test_choose_document_version_is_audited_and_keeps_original_rows(client):
    http, db = client
    for name, qty in [("first.csv", 24), ("second.csv", 30)]:
        http.post(
            "/api/v1/commercial/imports",
            files={"files": (name, csv_bytes(quantity=qty), "text/csv")},
        )
    data = http.get("/api/v1/commercial/workflow").json()
    assert len(data["orders"]) == 1
    order = data["orders"][0]
    assert order["status"] == "REVISAR"
    assert len(order["versions"]) == 2
    chosen = order["versions"][1]["records"][0]
    payload = dict(
        record_id=chosen["id"],
        reason="Original verificado por el usuario",
        revisions={
            r["id"]: r["revision"] for v in order["versions"] for r in v["records"]
        },
    )
    assert (
        http.post(
            "/api/v1/commercial/workflow/version", json={**payload, "revisions": {}}
        ).status_code
        == 409
    )
    result = http.post("/api/v1/commercial/workflow/version", json=payload)
    assert result.status_code == 200, result.text
    result = http.get("/api/v1/commercial/workflow").json()
    assert len(result["orders"]) == 1
    assert result["orders"][0]["quantity"] == 30
    assert result["orders"][0]["status"] == "NO FACTURADO"
    rows = list(db.scalars(select(CommercialRecord)))
    assert len(rows) == 2 and all(not r.superseded for r in rows)
    assert sorted(r.source["total_units"] for r in rows) == ["24", "30"]
    assert (
        len(
            list(
                db.scalars(
                    select(AuditLog).where(
                        AuditLog.action == "commercial.select_version"
                    )
                )
            )
        )
        == 2
    )
