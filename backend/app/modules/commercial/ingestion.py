"""Bounded multi-format ingestion; unsupported layouts remain reviewable evidence."""

import csv
import io
import mimetypes
import re
import zipfile
from datetime import date, datetime
from pathlib import PurePosixPath

from openpyxl import load_workbook

from app.modules.commercial.normalization import HEADER_MAP, KIND_MAP, identity

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_BATCH_BYTES = 60 * 1024 * 1024
MAX_FILES = 60
MAX_ROWS = 20000
EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tif",
    ".tiff",
}


def unpack(name, content):
    """Never extract an archive to disk; reject nested/oversized archives."""
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("Archivo supera 20 MB")
    if PurePosixPath(name).suffix.lower() != ".zip":
        return [(PurePosixPath(name).name, content)]
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = [
                entry
                for entry in archive.infolist()
                if not entry.is_dir() and not entry.filename.startswith("__MACOSX/")
            ]
            if (
                len(entries) > MAX_FILES
                or sum(entry.file_size for entry in entries) > MAX_BATCH_BYTES
            ):
                raise ValueError("ZIP supera 60 documentos o 60 MB descomprimidos")
            for entry in entries:
                path = PurePosixPath(entry.filename)
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or entry.file_size > MAX_FILE_BYTES
                ):
                    raise ValueError("ZIP contiene una ruta o tamaño no permitido")
                if entry.flag_bits & 1 or path.suffix.lower() == ".zip":
                    raise ValueError("ZIP cifrado o anidado no admitido")
            return [(PurePosixPath(e.filename).name, archive.read(e)) for e in entries]
    except zipfile.BadZipFile as exc:
        raise ValueError("ZIP inválido") from exc


def cell_value(value):
    return value.isoformat()[:10] if isinstance(value, (date, datetime)) else value


def table_records(rows, sheet, defaults=None):
    headers = None
    metadata = dict(defaults or {})
    if identity(sheet) in KIND_MAP:
        metadata["kind"] = KIND_MAP[identity(sheet)]
    result = []
    for row_index, values in enumerate(rows, 1):
        if row_index > MAX_ROWS + 100:
            raise ValueError("La hoja supera 20.000 filas; divide el archivo")
        values = list(values)
        detected = [HEADER_MAP.get(identity(v)) for v in values]
        recognized = {v for v in detected if v}
        if recognized & {"product_name", "code", "ean"} and recognized & {
            "quantity",
            "total_units",
            "boxes",
        }:
            headers = detected
            repeated = len([h for h in headers if h]) != len(recognized)
            if repeated:
                raise ValueError("Encabezados duplicados/ambiguos en la tabla")
            continue
        if headers is None:
            for index, value in enumerate(values):
                if identity(value) in ("sell in", "sell out"):
                    metadata["kind"] = KIND_MAP[identity(value)]
                parts = str(value or "").split(":", 1)
                key = HEADER_MAP.get(identity(parts[0]))
                if key in {
                    "kind",
                    "chain",
                    "date",
                    "currency",
                    "order_number",
                    "invoice_number",
                }:
                    candidate = (
                        parts[1].strip()
                        if len(parts) == 2
                        else values[index + 1]
                        if index + 1 < len(values)
                        else None
                    )
                    if candidate is not None and str(candidate).strip():
                        metadata[key] = cell_value(candidate)
            continue
        if not any(v is not None and str(v).strip() for v in values):
            continue
        raw = dict(metadata)
        for key, value in zip(headers, values):
            if key and value is not None and str(value).strip():
                raw[key] = cell_value(value)
        raw["source_location"] = f"{sheet}, fila {row_index}"
        raw["original_cells"] = [cell_value(v) for v in values]
        # Footer totals are document metadata, not product lines.
        if (
            re.fullmatch(
                r"(?:sub\s*total|total|iva|descuento)\s*:?",
                identity(raw.get("product_name")),
            )
            and not raw.get("code")
            and not raw.get("ean")
        ):
            continue
        result.append(raw)
        if len(result) > MAX_ROWS:
            raise ValueError("Archivo supera 20.000 registros")
    return result


def parse_file(name, content):
    suffix = PurePosixPath(name).suffix.lower()
    mime = mimetypes.guess_type(name)[0] or "application/octet-stream"
    if suffix not in EXTENSIONS:
        raise ValueError(
            "Formato no admitido. Usa PDF, XLSX, CSV, PNG, JPG, WEBP, TIFF o ZIP; convierte XLS a XLSX"
        )
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        records = table_records(csv.reader(io.StringIO(text), dialect), "CSV")
        return (
            mime,
            "csv",
            text,
            records
            or [
                {
                    "source_location": "CSV",
                    "extraction_issue": "No se reconocieron encabezados de producto y cantidad",
                }
            ],
            [],
        )
    if suffix == ".xlsx":
        # OOXML is itself a zip: check expansion before invoking the parser.
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(e.file_size for e in archive.infolist()) > MAX_BATCH_BYTES:
                raise ValueError("Excel supera 60 MB descomprimidos")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        records = []
        try:
            for sheet in workbook:
                sheet_records = table_records(
                    sheet.iter_rows(values_only=True), sheet.title
                )
                for row in sheet_records:
                    if any(
                        isinstance(v, str) and v.startswith("=")
                        for v in row.get("original_cells", [])
                    ):
                        row["extraction_issue"] = (
                            "La fila contiene fórmulas: exporta valores o confirma los datos desde el original"
                        )
                records.extend(sheet_records)
                if len(records) > MAX_ROWS:
                    raise ValueError("Excel supera 20.000 registros")
        finally:
            workbook.close()
        return (
            mime,
            "xlsx",
            None,
            records
            or [
                {
                    "extraction_issue": "No se reconocieron encabezados de producto y cantidad"
                }
            ],
            [],
        )
    import fitz
    from app.modules.purchase_orders.domain.document_extraction import (
        extract_document,
        recognized_header,
        split_purchase_orders,
    )
    from app.modules.purchase_orders.domain.table_extraction import (
        extract_known_order_text_rows,
    )

    if suffix == ".pdf":
        with fitz.open(stream=content, filetype="pdf") as pdf:
            if len(pdf) > 40:
                raise ValueError("PDF supera 40 páginas; divide el archivo")
    extracted = extract_document(content, mime, name)
    logical_orders = split_purchase_orders(extracted.text)
    header = (
        recognized_header(logical_orders[0])
        if len(logical_orders) > 1
        else (extracted.header or recognized_header(extracted.text))
    )
    from app.modules.commercial.invoice_extraction import (
        invoice_header,
        document_currency,
    )

    invoice = invoice_header(extracted.text)
    defaults = invoice or {
        "chain": header.get("chain_name"),
        "buyer_legal_name": header.get("buyer_legal_name"),
        "buyer_ruc": header.get("buyer_ruc"),
        "document_type": header.get("document_type"),
        "date": header.get("order_date"),
        "order_number": header.get("order_number"),
        "kind": "order"
        if header.get("order_number") or header.get("document_type")
        else None,
    }
    defaults["currency"] = document_currency(extracted.text)
    if not defaults.get("document_total"):
        total = re.search(r"(?im)^\s*Total\s+USD\s*:\s*([\d.,]+)", extracted.text)
        if total:
            defaults["document_total"] = total.group(1)
    records = []
    if len(logical_orders) > 1 and not invoice:
        for order_index, order_text in enumerate(logical_orders, 1):
            order_header = recognized_header(order_text)
            order_defaults = {
                "chain": order_header.get("chain_name"),
                "buyer_legal_name": order_header.get("buyer_legal_name"),
                "buyer_ruc": order_header.get("buyer_ruc"),
                "document_type": order_header.get("document_type"),
                "date": order_header.get("order_date"),
                "order_number": order_header.get("order_number"),
                "kind": "order",
            }
            for row in extract_known_order_text_rows(order_text):
                records.append(
                    {
                        **order_defaults,
                        "product_name": row.get("description"),
                        "code": row.get("supplier_reference") or row.get("chain_code"),
                        "quantity": row.get("quantity"),
                        "units_per_box": row.get("units_per_box"),
                        "unit": row.get("original_unit_type"),
                        "unit_price": row.get("unit_price"),
                        "amount": row.get("amount"),
                        "source_location": f"Orden {order_index}",
                        "source_text": row.get("raw"),
                    }
                )
    elif suffix == ".pdf" and (invoice or not header.get("confidence")):
        with fitz.open(stream=content, filetype="pdf") as pdf:
            for page_number, page in enumerate(pdf, 1):
                for table in page.find_tables().tables:
                    records.extend(
                        table_records(
                            table.extract(), f"Página {page_number}", defaults
                        )
                    )
    if not records:
        for row in extracted.table_rows:
            records.append(
                {
                    **defaults,
                    "product_name": row.get("description"),
                    "code": row.get("supplier_reference") or row.get("chain_code"),
                    "quantity": row.get("quantity"),
                    "units_per_box": row.get("units_per_box"),
                    "unit": row.get("original_unit_type"),
                    "unit_price": row.get("unit_price"),
                    "amount": row.get("amount"),
                    "source_location": f"Página {row.get('page', 1)}",
                    "source_text": row.get("raw"),
                }
            )
    if not records:
        records = [
            {
                **defaults,
                "source_location": "Documento",
                "extraction_issue": "No se detectó una tabla verificable; revisar el texto extraído y el original",
            }
        ]
    # Keep uncertainty once at document level instead of repeating it per product.
    recognized = recognized_header(extracted.text)
    document_issues = [] if invoice else list(recognized.get("pending_fields") or [])
    generic_ocr_review = extracted.method != "pdf_text" and not recognized.get(
        "confidence"
    )
    if generic_ocr_review:
        document_issues.append("Verificar cabecera extraída por OCR")
    if document_issues:
        warnings = [*extracted.warnings, *dict.fromkeys(document_issues)]
    else:
        warnings = list(extracted.warnings)
    if not invoice and (
        not recognized.get("chain_name")
        or not recognized.get("order_number")
        or generic_ocr_review
    ):
        for row in records:
            row["extraction_issue"] = (
                "Verificar los campos pendientes de la cabecera del documento"
            )
    return mime, extracted.method, extracted.text, records, warnings
