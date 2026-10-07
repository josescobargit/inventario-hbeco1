import hashlib
import io
import logging
import uuid
from datetime import date, timedelta
from typing import Annotated, Literal
from urllib.parse import quote

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.modules.audit.infrastructure.models import AuditLog
from app.modules.auth.api.dependencies import CurrentUser
from app.modules.commercial.assistant import plan_question
from app.modules.commercial.ingestion import (
    MAX_BATCH_BYTES,
    MAX_FILE_BYTES,
    MAX_FILES,
    parse_file,
    unpack,
)
from app.modules.commercial.models import (
    CommercialDocument,
    CommercialProductProfile,
    CommercialRecord,
)
from app.modules.commercial.normalization import HEADERS, KIND_MAP, identity
from app.modules.commercial.service import (
    analyze,
    catalog_context,
    executive_report,
    records,
)

from app.modules.commercial.processing import (
    PARSER_VERSION,
    processing_records,
    refresh_extraction,
)

router = APIRouter(prefix="/commercial", tags=["Control comercial"])
DB = Annotated[Session, Depends(get_db)]


class Filters(BaseModel):
    date_from: date | None = None
    date_to: date | None = None
    chain: str | None = Field(None, max_length=160)
    product_id: str | None = Field(None, max_length=40)
    line: str | None = Field(None, max_length=100)
    target: float = Field(95, ge=0, le=100)
    economic_threshold: float = Field(1000, ge=0, le=1e12)

    def parameters(self):
        if self.date_from and self.date_to and self.date_from > self.date_to:
            raise HTTPException(422, "La fecha inicial debe ser anterior a la final")
        return {
            "start": self.date_from.isoformat() if self.date_from else None,
            "end": self.date_to.isoformat() if self.date_to else None,
            "chain": self.chain,
            "product_id": self.product_id,
            "line": self.line,
            "target": self.target,
            "economic_threshold": self.economic_threshold,
        }


def projection(db, filters):
    result = analyze(records(db), **filters.parameters())
    result["executive_report"] = executive_report(result)
    return result


@router.get("/dashboard")
def dashboard(_user: CurrentUser, db: DB, filters: Annotated[Filters, Depends()]):
    return projection(db, filters)


@router.get("/workflow")
def workflow(_user: CurrentUser, db: DB):
    from app.modules.commercial.workflow import saved_workflow

    return saved_workflow(db)


class DocumentVersionChoice(BaseModel):
    record_id: uuid.UUID
    reason: str = Field(min_length=5, max_length=1000)
    revisions: dict[str, int]


@router.post("/workflow/version")
def choose_document_version(payload: DocumentVersionChoice, user: CurrentUser, db: DB):
    from app.modules.commercial.workflow import chain_name, reference

    chosen = db.get(CommercialRecord, payload.record_id)
    if not chosen or chosen.superseded:
        raise HTTPException(404, "Versión no encontrada")
    raw = {**chosen.source, **chosen.corrections}
    key = (
        KIND_MAP.get(identity(raw.get("kind"))),
        chain_name(raw.get("chain")),
        reference(
            raw.get("invoice_number")
            if KIND_MAP.get(identity(raw.get("kind"))) == "invoice"
            else raw.get("order_number")
        ),
    )
    if not all(key):
        raise HTTPException(
            422, "Completa tipo, cadena y número antes de elegir una versión"
        )
    related = []
    for row in db.scalars(
        select(CommercialRecord)
        .where(CommercialRecord.superseded.is_(False))
        .with_for_update()
    ):
        value = {**row.source, **row.corrections}
        if (
            KIND_MAP.get(identity(value.get("kind"))),
            chain_name(value.get("chain")),
            reference(
                value.get("invoice_number")
                if KIND_MAP.get(identity(value.get("kind"))) == "invoice"
                else value.get("order_number")
            ),
        ) == key:
            related.append(row)
    if {str(r.id): r.revision for r in related} != payload.revisions:
        raise HTTPException(409, "Las versiones cambiaron. Actualiza antes de elegir.")
    for row in related:
        previous = dict(row.corrections)
        row.corrections = {
            **previous,
            "status": "activo" if row.document_id == chosen.document_id else "excluido",
        }
        row.revision += 1
        db.add(
            AuditLog(
                actor_user_id=user.id,
                action="commercial.select_version",
                entity_type="commercial_record",
                entity_id=str(row.id),
                previous_value=previous,
                new_value={"corrections": row.corrections},
                reason=payload.reason,
            )
        )
    db.commit()
    return {"status": "OK"}


@router.get("/documents")
def documents(_user: CurrentUser, db: DB):
    return [
        {
            "id": str(d.id),
            "filename": d.filename,
            "sha256": d.sha256,
            "method": d.method,
            "parser_version": d.parser_version,
            "needs_reprocessing": d.parser_version != PARSER_VERSION,
            "warnings": d.warnings,
            "created_at": d.created_at,
        }
        for d in db.scalars(
            select(CommercialDocument).order_by(CommercialDocument.created_at.desc())
        )
    ]


@router.get("/processing")
def processing(_user: CurrentUser, db: DB):
    return {"records": processing_records(db)}


@router.post("/documents/{document_id}/reprocess")
def reprocess_document(document_id: uuid.UUID, user: CurrentUser, db: DB):
    document = db.scalar(
        select(CommercialDocument)
        .where(CommercialDocument.id == document_id)
        .with_for_update()
    )
    if document is None:
        raise HTTPException(404, "Documento no encontrado")
    try:
        count = refresh_extraction(
            db, document, parse_file(document.filename, document.content), user.id
        )
        db.commit()
    except ValueError as exc:
        db.rollback()
        raise HTTPException(422, str(exc))
    return {
        "document_id": str(document.id),
        "filename": document.filename,
        "status": "processed",
        "rows": count,
    }


@router.get("/documents/{document_id}/content")
def content(document_id: uuid.UUID, _user: CurrentUser, db: DB):
    document = db.get(CommercialDocument, document_id)
    if document is None:
        raise HTTPException(404, "Documento no encontrado")
    return Response(
        document.content,
        media_type=document.content_type,
        headers={
            "Content-Disposition": f"attachment; filename*=UTF-8''{quote(document.filename, safe='')}",
            "X-Content-Type-Options": "nosniff",
        },
    )


@router.get("/documents/{document_id}/text")
def document_text(document_id: uuid.UUID, _user: CurrentUser, db: DB):
    document = db.get(CommercialDocument, document_id)
    if document is None:
        raise HTTPException(404, "Documento no encontrado")
    return {"text": document.extracted_text, "warnings": document.warnings}


@router.post("/imports")
def import_files(user: CurrentUser, db: DB, files: Annotated[list[UploadFile], File()]):
    if len(files) > MAX_FILES:
        raise HTTPException(413, "Máximo 60 archivos por carga")
    inputs = []
    total_bytes = 0
    for file in files:
        payload = file.file.read(MAX_FILE_BYTES + 1)
        try:
            expanded = unpack(file.filename or "archivo", payload)
        except ValueError as exc:
            raise HTTPException(422, str(exc)) from exc
        total_bytes += sum(len(value) for _, value in expanded)
        inputs.extend(expanded)
        if total_bytes > MAX_BATCH_BYTES or len(inputs) > MAX_FILES:
            raise HTTPException(
                413, "La carga supera 60 archivos o 60 MB descomprimidos"
            )
    results = []
    for filename, payload in inputs:
        digest = hashlib.sha256(payload).hexdigest()
        previous = db.scalar(
            select(CommercialDocument)
            .where(CommercialDocument.sha256 == digest)
            .with_for_update()
        )
        if previous:
            if previous.parser_version != PARSER_VERSION:
                try:
                    count = refresh_extraction(
                        db, previous, parse_file(filename, payload), user.id
                    )
                    db.commit()
                    results.append(
                        {
                            "filename": filename,
                            "status": "processed",
                            "document_id": str(previous.id),
                            "rows": count,
                            "message": "Documento actualizado con el motor vigente",
                        }
                    )
                except Exception as exc:
                    db.rollback()
                    logging.getLogger(__name__).exception(
                        "Reprocessing failed for %s", filename
                    )
                    results.append(
                        {
                            "filename": filename,
                            "status": "error",
                            "message": str(exc)
                            if isinstance(exc, ValueError)
                            else "No se pudo reprocesar el original",
                        }
                    )
            else:
                results.append(
                    {
                        "filename": filename,
                        "status": "duplicate",
                        "document_id": str(previous.id),
                        "message": "Documento ya procesado; se muestran sus resultados",
                    }
                )
            continue
        try:
            mime, method, text, source_rows, warnings = parse_file(filename, payload)
            with db.begin_nested():
                document = CommercialDocument(
                    filename=filename[:255],
                    sha256=digest,
                    content_type=mime,
                    content=payload,
                    extracted_text=text,
                    method=method,
                    parser_version=PARSER_VERSION,
                    warnings=warnings,
                    created_by_user_id=user.id,
                )
                db.add(document)
                db.flush()
                for index, raw in enumerate(source_rows, 1):
                    db.add(
                        CommercialRecord(
                            document_id=document.id,
                            row_number=index,
                            source=raw,
                            corrections={},
                        )
                    )
                db.add(
                    AuditLog(
                        actor_user_id=user.id,
                        action="commercial.import",
                        entity_type="commercial_document",
                        entity_id=str(document.id),
                        new_value={
                            "filename": filename,
                            "sha256": digest,
                            "rows": len(source_rows),
                        },
                    )
                )
            db.commit()
            results.append(
                {
                    "filename": filename,
                    "status": "processed",
                    "document_id": str(document.id),
                    "rows": len(source_rows),
                    "message": "Procesado; consulta las alertas y filas por revisar",
                }
            )
        except IntegrityError:
            db.rollback()
            results.append(
                {
                    "filename": filename,
                    "status": "duplicate",
                    "message": "Importación concurrente ya registrada; actualiza la vista",
                }
            )
        except Exception as exc:
            db.rollback()
            logging.getLogger(__name__).exception(
                "Commercial import failed for %s", filename
            )
            results.append(
                {
                    "filename": filename,
                    "status": "error",
                    "message": str(exc)
                    if isinstance(exc, ValueError)
                    else "No se pudo extraer el archivo; revisa formato, cifrado o disponibilidad de OCR",
                }
            )
    return {"files": results}


class Correction(BaseModel):
    revision: int = Field(ge=1)
    reason: str = Field(min_length=5, max_length=2000)
    values: dict


@router.patch("/records/{record_id}")
def correct_record(record_id: uuid.UUID, body: Correction, user: CurrentUser, db: DB):
    row = db.scalar(
        select(CommercialRecord)
        .where(CommercialRecord.id == record_id)
        .with_for_update()
    )
    if row is None or row.superseded:
        raise HTTPException(404, "Registro no encontrado o extracción reemplazada")
    if row.revision != body.revision:
        raise HTTPException(409, "El registro cambió; actualiza antes de corregir")
    allowed = set(HEADERS) | {"product_id", "extraction_confirmed"}
    if set(body.values) - allowed or any(
        isinstance(v, (list, dict)) or isinstance(v, str) and len(v) > 500
        for v in body.values.values()
    ):
        raise HTTPException(422, "Campos de corrección inválidos")
    if "extraction_confirmed" in body.values and not isinstance(
        body.values["extraction_confirmed"], bool
    ):
        raise HTTPException(422, "La confirmación de extracción debe ser booleana")
    previous = dict(row.corrections)
    row.corrections = {**row.corrections, **body.values}
    row.revision += 1
    db.add(
        AuditLog(
            actor_user_id=user.id,
            action="commercial.correct",
            entity_type="commercial_record",
            entity_id=str(row.id),
            reason=body.reason,
            previous_value=previous,
            new_value=row.corrections,
        )
    )
    db.commit()
    return {"id": str(row.id), "revision": row.revision}


class DocumentDefaults(BaseModel):
    reason: str = Field(min_length=5, max_length=2000)
    values: dict


@router.patch("/documents/{document_id}/missing-header")
def complete_header(
    document_id: uuid.UUID, body: DocumentDefaults, user: CurrentUser, db: DB
):
    allowed = {"kind", "chain", "date", "currency", "order_number", "invoice_number"}
    if (
        not body.values
        or set(body.values) - allowed
        or any(
            not isinstance(value, str) or not value.strip() or len(value) > 160
            for value in body.values.values()
        )
    ):
        raise HTTPException(
            422, "Completa únicamente datos de cabecera presentes en el original"
        )
    rows = list(
        db.scalars(
            select(CommercialRecord)
            .where(
                CommercialRecord.document_id == document_id,
                CommercialRecord.superseded.is_(False),
            )
            .with_for_update()
        )
    )
    if not rows:
        raise HTTPException(404, "Documento no encontrado")
    changed = 0
    for row in rows:
        merged = {**row.source, **row.corrections}
        changes = {
            key: value for key, value in body.values.items() if not merged.get(key)
        }
        if not changes:
            continue
        previous = dict(row.corrections)
        row.corrections = {**previous, **changes}
        row.revision += 1
        changed += 1
        db.add(
            AuditLog(
                actor_user_id=user.id,
                action="commercial.complete_header",
                entity_type="commercial_record",
                entity_id=str(row.id),
                reason=body.reason,
                previous_value=previous,
                new_value=row.corrections,
            )
        )
    db.commit()
    return {"updated_records": changed}


class DocumentCorrection(DocumentDefaults):
    revisions: dict[str, int]


@router.patch("/documents/{document_id}/header")
def correct_document_header(
    document_id: uuid.UUID, body: DocumentCorrection, user: CurrentUser, db: DB
):
    allowed = {
        "kind",
        "chain",
        "date",
        "currency",
        "order_number",
        "invoice_number",
        "extraction_confirmed",
        "status",
    }
    if (
        not body.values
        or set(body.values) - allowed
        or any(
            not isinstance(value, str) or not value.strip() or len(value) > 160
            for key, value in body.values.items()
            if key != "extraction_confirmed"
        )
    ):
        raise HTTPException(422, "Selecciona los datos de cabecera que deseas corregir")
    if "extraction_confirmed" in body.values and not isinstance(
        body.values["extraction_confirmed"], bool
    ):
        raise HTTPException(422, "La verificación debe ser booleana")
    if "date" in body.values:
        from datetime import date

        try:
            date.fromisoformat(body.values["date"])
        except ValueError:
            raise HTTPException(422, "La fecha debe ser válida (AAAA-MM-DD)")
    if "kind" in body.values and body.values["kind"] not in {
        "order",
        "invoice",
        "sell_in",
        "sell_out",
    }:
        raise HTTPException(422, "Selecciona un tipo de documento válido")
    if "status" in body.values and body.values["status"] not in {
        "activo",
        "excluido",
        "anulada",
        "reemplazada",
    }:
        raise HTTPException(422, "Selecciona un estado administrativo válido")
    rows = list(
        db.scalars(
            select(CommercialRecord)
            .where(
                CommercialRecord.document_id == document_id,
                CommercialRecord.superseded.is_(False),
            )
            .with_for_update()
        )
    )
    if not rows:
        raise HTTPException(404, "Documento no encontrado")
    if set(body.revisions) != {str(row.id) for row in rows} or any(
        body.revisions.get(str(row.id)) != row.revision for row in rows
    ):
        raise HTTPException(409, "El documento cambió; actualiza antes de corregir")
    for row in rows:
        previous = dict(row.corrections)
        row.corrections = {**previous, **body.values}
        row.revision += 1
        db.add(
            AuditLog(
                actor_user_id=user.id,
                action="commercial.correct_document_header",
                entity_type="commercial_record",
                entity_id=str(row.id),
                reason=body.reason,
                previous_value=previous,
                new_value=row.corrections,
            )
        )
    db.commit()
    return {"updated_records": len(rows)}


@router.get("/products")
def products(_user: CurrentUser, db: DB):
    catalog, profiles, aliases = catalog_context(db)
    return [
        {
            "id": str(p.id),
            "sku": p.sku,
            "name": p.name,
            "description": p.description,
            "ean13": p.barcode,
            "units_per_box": p.units_per_box,
            "category": p.category,
            **profiles.get(str(p.id), {}),
            "chain_aliases": [
                {"chain": a.chain_name, "name": a.source_text, "code": a.detected_code}
                for a in aliases
                if a.product_id == p.id
            ],
        }
        for p in catalog
    ]


class ProfileInput(BaseModel):
    ean14: str | None = Field(None, pattern=r"^\d{14}$")
    presentation: str | None = Field(None, max_length=160)
    content: str | None = Field(None, max_length=160)
    line: str | None = Field(None, max_length=100)
    aliases: list[str] = Field(default_factory=list, max_length=100)
    reason: str = Field(min_length=5, max_length=2000)


@router.put("/products/{product_id}")
def profile(product_id: uuid.UUID, body: ProfileInput, user: CurrentUser, db: DB):
    if user.role.code != "principal":
        raise HTTPException(
            403, "Solo el administrador confirma equivalencias globales"
        )
    catalog, _, _ = catalog_context(db)
    if product_id not in {p.id for p in catalog}:
        raise HTTPException(404, "Producto no encontrado")
    if any(not alias.strip() or len(alias) > 300 for alias in body.aliases):
        raise HTTPException(422, "Equivalencias vacías o demasiado largas")
    if body.ean14:
        digits = [int(d) for d in body.ean14]
        if (
            sum(d * (3 if i % 2 == 0 else 1) for i, d in enumerate(digits[:-1])) % 10
            != (-digits[-1]) % 10
        ):
            raise HTTPException(422, "Dígito verificador EAN 14 inválido")
    row = db.get(CommercialProductProfile, product_id)
    previous = (
        {
            key: getattr(row, key)
            for key in ("ean14", "presentation", "content", "line", "aliases")
        }
        if row
        else None
    )
    if row is None:
        row = CommercialProductProfile(product_id=product_id)
        db.add(row)
    for key, value in body.model_dump(exclude={"reason"}).items():
        setattr(row, key, value)
    db.add(
        AuditLog(
            actor_user_id=user.id,
            action="commercial.product_profile",
            entity_type="product",
            entity_id=str(product_id),
            reason=body.reason,
            previous_value=previous,
            new_value=body.model_dump(exclude={"reason"}),
        )
    )
    db.commit()
    return {"id": str(product_id)}


class CompareInput(Filters):
    previous_from: date
    previous_to: date


@router.post("/compare")
def compare(body: CompareInput, _user: CurrentUser, db: DB):
    if body.previous_from > body.previous_to:
        raise HTTPException(422, "Período anterior inválido")
    source = records(db)
    current = analyze(source, **body.parameters())
    parameters = body.parameters()
    parameters.update(
        start=body.previous_from.isoformat(), end=body.previous_to.isoformat()
    )
    previous = analyze(source, **parameters)
    changes = []
    for key in (
        "ordered",
        "invoiced",
        "missing",
        "ordered_amount",
        "invoiced_amount",
        "unbilled_amount",
        "sell_in_units",
        "sell_out_units",
    ):
        a, b = current["summary"].get(key), previous["summary"].get(key)
        changes.append(
            {
                "metric": key,
                "current": a,
                "previous": b,
                "difference": a - b if a is not None and b is not None else None,
                "variation": round((a - b) / abs(b) * 100, 2)
                if a is not None and b
                else None,
            }
        )
    sales_changes = []
    for kind, chain, pid in {
        (s["kind"], s["chain"], s["product_id"])
        for s in [*current["sales"], *previous["sales"]]
    }:
        a = next(
            (
                s
                for s in current["sales"]
                if (s["kind"], s["chain"], s["product_id"]) == (kind, chain, pid)
            ),
            None,
        )
        b = next(
            (
                s
                for s in previous["sales"]
                if (s["kind"], s["chain"], s["product_id"]) == (kind, chain, pid)
            ),
            None,
        )
        sales_changes.append(
            {
                "kind": kind,
                "chain": chain,
                "product_name": (a or b)["product_name"],
                "current": a["units"] if a else None,
                "previous": b["units"] if b else None,
                "variation": round((a["units"] - b["units"]) / b["units"] * 100, 2)
                if a and b and b["units"]
                else None,
            }
        )
    return {
        "current": current["filters"],
        "previous": previous["filters"],
        "changes": changes,
        "sales_changes": sales_changes,
        "caveat": "Períodos sin registros no prueban cero ventas; verifica cobertura antes de interpretar crecimiento o caída.",
    }


class Question(Filters):
    question: str = Field(min_length=3, max_length=1500)


@router.post("/assistant")
def assistant(body: Question, user: CurrentUser, db: DB):
    from zoneinfo import ZoneInfo
    from datetime import datetime

    today = datetime.now(ZoneInfo("America/Guayaquil")).date()
    plan, mode = plan_question(
        body.question, body.model_dump(exclude={"question"}), today
    )
    if plan:
        if plan.date_from:
            body.date_from = plan.date_from
        if plan.date_to:
            body.date_to = plan.date_to
        if plan.chain:
            body.chain = plan.chain
    intents = {
        "summary": "resumen",
        "alerts": "anomalias",
        "missing": "productos no facturados",
        "sell_in": "sell in",
        "sell_out": "sell out",
        "sell_comparison": "sell in y sell out",
        "weekly_comparison": "compara semana",
        "unsupported": "consulta no soportada",
    }
    query = intents[plan.intent] if plan else identity(body.question)
    if "semana" in query and not body.date_from:
        from zoneinfo import ZoneInfo
        from datetime import datetime

        today = datetime.now(ZoneInfo("America/Guayaquil")).date()
        body.date_from = today - timedelta(days=today.weekday())
        body.date_to = today
    data = projection(db, body)
    response = []
    if "compara" in query and "semana" in query:
        start = body.date_from or date.today()
        end = body.date_to or start + timedelta(days=6)
        comparison = compare(
            CompareInput(
                **body.model_dump(exclude={"question"}),
                previous_from=start - timedelta(days=7),
                previous_to=end - timedelta(days=7),
            ),
            user,
            db,
        )
        response = [
            f"{c['metric']}: actual {c['current'] if c['current'] is not None else 'sin dato'}, anterior {c['previous'] if c['previous'] is not None else 'sin dato'}; variación {str(c['variation']) + '%' if c['variation'] is not None else 'no calculable'}."
            for c in comparison["changes"]
        ]
    elif "sell in" in query and "sell out" in query:
        response = [
            f"{s['chain']} · {s['product_name']}: Sell In {s['sell_in']}, Sell Out {s['sell_out']}. Hipótesis: {s['hypothesis']}"
            for s in data["sell_comparison"][:20]
        ]
    elif "sell in" in query or "sell out" in query:
        kind = "sell_in" if "sell in" in query else "sell_out"
        response = [
            f"{s['chain']} · {s['product_name']}: {s['units']:g} unidades; USD {s['amount'] if s['amount'] is not None else 'sin dato'}; participación {s['share']}%."
            for s in sorted(data["sales"], key=lambda x: x["units"], reverse=True)
            if s["kind"] == kind
        ][:20]
    elif any(word in query for word in ("anomalia", "revisar", "problema")):
        response = [f"{a['level']}: {a['message']}" for a in data["alerts"][:20]]
    elif "producto" in query and any(
        word in query for word in ("falt", "no se factur", "no factur")
    ):
        response = [
            f"{r['chain']} · OC {r['order_number']} · {r['product_name']}: {r['missing']:g} unidades faltantes ({r['status']})."
            for r in data["rows"]
            if r["missing"]
        ][:20]
    elif any(
        word in query
        for word in (
            "analiza",
            "resumen",
            "gerencia",
            "dinero",
            "facturar",
            "pendiente",
            "semana",
        )
    ):
        response = data["executive_report"]
    else:
        response = [
            "Puedo consultar el resumen, dinero pendiente, faltantes, anomalías, Sell In, Sell Out y comparar esta semana con la anterior. No encuentro una consulta verificable para esa pregunta."
        ]
    return {
        "answer": response
        or [
            "No hay registros suficientes para responder con los filtros seleccionados."
        ],
        "mode": mode,
        "filters": data["filters"],
        "basis": data["basis"],
        "sources": [
            {
                "id": r["id"],
                "filename": r["filename"],
                "location": r["source_location"],
                "url": r["source_url"],
            }
            for r in data["records"]
        ],
        "limitations": "Respuestas por consultas verificables, sin generación libre. Hipótesis no prueban existencias ni causas de una anomalía.",
    }


EXPORT_COLUMNS = [
    ("chain", "CADENA"),
    ("order_number", "OC"),
    ("sku", "SKU"),
    ("product_name", "PRODUCTO"),
    ("ordered", "PEDIDO OC"),
    ("invoiced", "FACTURADO"),
    ("difference", "DIFERENCIA"),
    ("fulfillment", "% ENTREGA"),
    ("missing", "FALTANTE"),
    ("ordered_amount", "DINERO PEDIDO USD"),
    ("invoiced_amount", "DINERO FACTURADO USD"),
    ("unbilled_amount", "DINERO NO FACTURADO USD"),
    ("status", "ESTADO"),
]


@router.get("/export")
def export(
    _user: CurrentUser,
    db: DB,
    filters: Annotated[Filters, Depends()],
    format: Literal["xlsx", "pdf"] = "xlsx",
    view: Literal["general", "chain", "detail", "missing", "report"] = "general",
):
    data = projection(db, filters)
    if format == "pdf":
        import fitz

        pdf = fitz.open()
        lines = [
            "CONTROL COMERCIAL · REPORTE GERENCIAL",
            f"Período: {filters.date_from or 'histórico'} a {filters.date_to or 'último registro'}",
            f"Cadena: {filters.chain or 'todas'}",
            *data["executive_report"],
            "BASE DE CÁLCULO",
            data["basis"],
        ]
        import textwrap

        rendered = [
            part for line in lines for part in [*textwrap.wrap(line, width=88), ""]
        ]
        for start in range(0, len(rendered), 45):
            page = pdf.new_page(width=595, height=842)
            page.insert_text(
                (42, 48),
                "\n".join(rendered[start : start + 45]),
                fontsize=10,
                lineheight=1.5,
            )
        result = pdf.tobytes()
        pdf.close()
        return Response(
            result,
            media_type="application/pdf",
            headers={
                "Content-Disposition": 'attachment; filename="reporte-comercial.pdf"'
            },
        )
    workbook = Workbook()
    workbook.remove(workbook.active)

    def sheet(name, columns, items):
        ws = workbook.create_sheet(name)
        ws.append([label for _, label in columns])
        for item in items:
            ws.append([item.get(key) for key, _ in columns])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="164E63")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = (
                        "s"  # Untrusted strings never become Excel formulas.
                    )
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(
                48, max(16, len(str(column[0].value or "")) + 3)
            )

    detail_rows = [r for r in data["rows"] if view != "missing" or r["missing"] > 0]
    sheet("OC vs factura", EXPORT_COLUMNS, detail_rows)
    sheet("Por cadena", EXPORT_COLUMNS, data["chains"])
    sheet("Cadena y producto", EXPORT_COLUMNS, data["by_chain_product"])
    sheet(
        "Sell In y Out",
        [
            (k, k.upper())
            for k in (
                "kind",
                "chain",
                "product_name",
                "line",
                "units",
                "amount",
                "share",
            )
        ],
        data["sales"],
    )
    sheet(
        "Trazabilidad",
        [
            (k, k.upper())
            for k in (
                "id",
                "filename",
                "source_location",
                "kind",
                "chain",
                "date",
                "order_number",
                "invoice_number",
                "product_name",
                "units",
                "conversion",
                "amount",
                "source_url",
                "review",
            )
        ],
        [
            {
                **r,
                "amount": float(r["amount"]) if r["amount"] is not None else None,
                "review": "; ".join([*r["issues"], *r["monetary_issues"]]),
            }
            for r in data["records"]
        ],
    )
    sheet(
        "Alertas",
        [("level", "PRIORIDAD"), ("message", "DETALLE"), ("refs", "REGISTROS")],
        [{**a, "refs": ", ".join(a["record_ids"])} for a in data["alerts"]],
    )
    sheet(
        "Reporte gerencial",
        [("text", "RESUMEN EJECUTIVO")],
        [
            {"text": t}
            for t in [
                *data["executive_report"],
                data["basis"],
                f"Filtros: {data['filters']}",
            ]
        ],
    )
    stream = io.BytesIO()
    workbook.save(stream)
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="comercial-{view}.xlsx"'
        },
    )
